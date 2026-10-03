"""PowerPoint (.pptx) parser. Synthetic decks (always run) plus checks
against the course samples (skipped when they aren't present)."""

import pymupdf as fitz
import pptx
import pytest
from conftest import SAMPLES, find, parse_sample
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Pt

from notes2gdoc.parsers import parse_file

CAPITAL = SAMPLES / "Lecture - Capital Allocation.pptx"
MULTIFACTOR = SAMPLES / "Lecture - Multifactor Models.pptx"
MGMT = SAMPLES / "MGMT102_2627_04.pptx"


def _png() -> bytes:
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 200, 120), False)
    pix.clear_with(150)
    return pix.tobytes("png")


@pytest.fixture(scope="module")
def deck(tmp_path_factory):
    prs = pptx.Presentation()
    title = prs.slides.add_slide(prs.slide_layouts[0])
    title.shapes.title.text = "Corporate Finance"
    title.placeholders[1].text = "Week 1"

    for _ in range(2):  # two slides with the same title -> one heading
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = "Risk and Return"
        tf = s.placeholders[1].text_frame
        tf.text = "Top point"
        sub = tf.add_paragraph()
        sub.text = "Sub point"
        sub.level = 1
        p = tf.add_paragraph()
        r = p.add_run()
        r.text = "Coloured "
        r.font.color.rgb = RGBColor(0x00, 0x70, 0xC0)
        r2 = p.add_run()
        r2.text = "3"
        r3 = p.add_run()
        r3.text = "rd"
        r3._r.get_or_add_rPr().set("baseline", "30000")
        r4 = p.add_run()
        r4.text = " s"
        r4.font.name = "Symbol"  # displays as σ

    sec = prs.slides.add_slide(prs.slide_layouts[2])  # "Section Header"
    sec.shapes.title.text = "Part Two"

    # Title in a plain text box (no title placeholder), plus a table
    s = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    tb = s.shapes.add_textbox(Pt(30), Pt(15), Pt(300), Pt(40))
    tb.text_frame.text = "Comparison"
    gt = s.shapes.add_table(2, 2, Pt(40), Pt(100), Pt(500), Pt(100)).table
    gt.cell(0, 0).merge(gt.cell(0, 1))
    gt.cell(0, 0).text = "Header"
    gt.cell(1, 0).text = "Left"
    gt.cell(1, 1).text = "Right"
    gt.cell(1, 0).fill.solid()
    gt.cell(1, 0).fill.fore_color.rgb = RGBColor(0xFF, 0xD9, 0x66)

    # A diagram: two boxes joined by an arrow connector; plus a picture
    s = prs.slides.add_slide(prs.slide_layouts[5])  # title only
    s.shapes.title.text = "Process"
    a = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(60), Pt(150), Pt(150), Pt(60))
    a.text_frame.text = "Step A"
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(400), Pt(150), Pt(150), Pt(60))
    b.text_frame.text = "Step B"
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(210), Pt(180), Pt(400), Pt(180))
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd", {"type": "triangle"})
    ln.append(tail)

    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Exhibit"
    path = tmp_path_factory.mktemp("pptx") / "deck.pptx"
    img = path.parent / "pic.png"
    img.write_bytes(_png())
    s.shapes.add_picture(str(img), Pt(100), Pt(120), Pt(300), Pt(180))
    prs.save(path)
    return parse_file(path)


def test_title_slide_unticked(deck):
    first = [b for b in deck.blocks if b.page == 1]
    assert [b.text for b in first] == ["Corporate Finance", "Week 1"]
    assert not any(b.selected for b in first)


def test_titles_merged_and_bullets(deck):
    titles = [b.text for b in deck.blocks if b.kind == "heading" and b.style_key == "slide_title"]
    assert titles.count("Risk and Return") == 1
    slide2 = [(b.kind, b.level, b.text) for b in deck.blocks if b.page == 2]
    assert slide2[1:3] == [("bullet", 0, "Top point"), ("bullet", 1, "Sub point")]
    styled = next(b for b in deck.blocks if b.text.startswith("Coloured"))
    assert styled.runs[0].color == "#0070C0"
    assert [r.text for r in styled.runs if r.superscript] == ["rd"]
    assert styled.text.endswith("σ")


def test_section_header(deck):
    h = find(deck, "Part Two", "heading")
    assert h.style_key == "section_title"


def test_textbox_title_and_table(deck):
    assert find(deck, "Comparison", "heading").style_key == "slide_title"
    t = next(b for b in deck.blocks if b.kind == "table").table
    assert t.cell(0, 0).colspan == 2 and t.cell(0, 0).text == "Header"
    assert t.cell(1, 0).background == "#FFD966"


def test_diagram_redrawn_and_picture(deck):
    pics = [b for b in deck.blocks if b.kind == "image"]
    assert len(pics) == 2
    diagram = next(b for b in pics if "redrawn" in b.note)
    assert "Step A" in diagram.image.alt and "Step B" in diagram.image.alt
    assert diagram.image.png.startswith(b"\x89PNG")
    # The boxes' text isn't repeated as paragraphs
    assert not any(b.text == "Step A" for b in deck.blocks if b.kind != "image")


# --------------------------------------------------------------------------- #
# Course samples (local only)
# --------------------------------------------------------------------------- #

def test_capital_allocation_sample():
    doc = parse_sample(CAPITAL)
    assert find(doc, "Overview", "heading").style_key == "slide_title"          # text-box title
    order = [b.text for b in doc.blocks if b.page == 5]
    assert order.index("High") < order.index("Less willing to take on more risks.") < order.index("Low")
    wrapped = find(doc, "Would not invest in zero risk premium securities")
    nxt = doc.blocks[doc.blocks.index(wrapped) + 1]
    assert nxt.text == "“fair games”." and nxt.kind == "paragraph" and nxt.level == wrapped.level + 1
    assert "→" in find(doc, "Expected profits is zero").text
    assert any(b.kind == "image" and "redrawn" in b.note and b.page == 21 for b in doc.blocks)
    assert not any(b.kind == "image" and b.page == 46 for b in doc.blocks)    # boxed title isn't a diagram
    equations = [b for b in doc.blocks if b.kind == "image" and "equation" in b.note]
    assert len(equations) == 11 and not doc.warnings


def test_multifactor_sample():
    doc = parse_sample(MULTIFACTOR)
    # Slide 4: β typed as "b" in the Symbol font; lines indented with spaces
    slide4 = [(b.kind, b.level, b.text) for b in doc.blocks if b.page == 4]
    assert ("paragraph", 1, "= E(Ri ) + βi GDP GDP + βi IR IR + ei") in slide4
    defs = [x for x in slide4 if "sensitivity of stock" in x[2]]
    assert [x[1] for x in defs] == [1, 1] and defs[0][2].startswith("βi GDP")
    where = next(x for x in slide4 if x[2].startswith("where"))
    assert where[1] == 0
    diagram = next(b for b in doc.blocks if b.kind == "image" and b.page == 3)
    assert "Security A" in diagram.image.alt and "Factor 3" in diagram.image.alt


def test_mgmt_sample():
    doc = parse_sample(MGMT)
    assert find(doc, "Learning Objectives", "heading")
    objectives = [b for b in doc.blocks if b.page == 2 and b.kind == "bullet"]
    assert objectives[0].text.startswith("Understand core competencies") and objectives[0].numbered
    assert len(objectives) == 7
    assert sum(1 for b in doc.blocks if b.kind == "image") >= 10
    assert not any("text alternative" in b.text.lower() for b in doc.blocks)
    assert not any("all rights reserved" in b.text.lower() for b in doc.blocks if b.selected)


def test_symbol_font_is_only_for_symbols():
    from notes2gdoc.parsers.symbols import fix_symbols

    # PowerPoint's separate "sym" font (Wingdings) doesn't change ordinary letters
    assert fix_symbols("Stocks with above", "Calibri", "Wingdings") == "Stocks with above"
    # ...but a run whose own font is Symbol does ("b" shows as beta)
    assert fix_symbols("b", "Symbol") == "β"


def test_samples_letters_and_template_bullets():
    multi = parse_sample(MULTIFACTOR)
    assert find(multi, "Stocks with above characteristics").kind == "bullet"
    assert not any("■" in b.text or "❖" in b.text for b in multi.blocks)
    # MGMT102's template has no bullet on top-level lines, only on the level below
    mgmt = parse_sample(MGMT)
    assert find(mgmt, "Exhibit 4.2").kind == "paragraph"
    assert find(mgmt, "Adjust along with the external environment").kind == "bullet"
