"""Word (.docx) parser. Synthetic documents (always run) plus checks against
the course samples (skipped when they aren't present)."""

import docx
import pymupdf as fitz
import pytest
from conftest import SAMPLES, find, parse_sample
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from notes2gdoc.parsers import ParseError, parse_file

BANKING = SAMPLES / "Banking & Fundraising Notes.docx"
CIV = SAMPLES / "CIV Exams.docx"


def _png() -> bytes:
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 40, 20), False)
    pix.clear_with(180)
    return pix.tobytes("png")


def _set_level(paragraph, level: int):
    """Put a list paragraph at a nesting level (Word's ilvl)."""
    numPr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    numPr.get_or_add_ilvl().val = level


@pytest.fixture(scope="module")
def sample_docx(tmp_path_factory):
    d = docx.Document()
    d.add_paragraph("My Notes", style="Title")
    d.add_heading("1. Corporate Financing", level=1)
    p = d.add_paragraph()
    r = p.add_run("(1) Purpose")
    r.bold = True
    r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    p.add_run(": companies need ")
    r = p.add_run("working capital")
    r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    p.add_run(" (3")
    r = p.add_run("rd")
    r.font.superscript = True
    p.add_run(" schedule) per ")
    r = p.add_run("Tung Hui")
    r.italic = True
    d.add_paragraph("")
    d.add_paragraph("")  # several blank lines collapse into one
    d.add_paragraph("Top bullet", style="List Bullet")
    sub = d.add_paragraph("Sub bullet", style="List Bullet")
    _set_level(sub, 1)
    d.add_paragraph("First step", style="List Number")
    d.add_paragraph("Second step", style="List Number")
    d.add_heading("A. Sub heading", level=2)
    line = d.add_paragraph("Line one")
    line.add_run().add_break()
    line.add_run("line two")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).merge(t.cell(0, 1)).text = "Merged header"
    t.cell(1, 0).text = "Left"
    t.cell(1, 1).text = "Right"
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "9CC2E5")
    t.cell(1, 0)._tc.get_or_add_tcPr().append(shd)
    d.add_paragraph()
    path = tmp_path_factory.mktemp("docx") / "sample.docx"
    img = path.parent / "pic.png"
    img.write_bytes(_png())
    d.add_picture(str(img), width=Pt(200))
    d.save(path)
    return parse_file(path)


def test_title_unticked_and_headings(sample_docx):
    b = sample_docx.blocks
    assert b[0].text == "My Notes" and not b[0].selected
    h1 = find(sample_docx, "1. Corporate Financing")
    assert (h1.kind, h1.style_key) == ("heading", "word_h1")
    assert find(sample_docx, "A. Sub heading").style_key == "word_h2"


def test_inline_formatting(sample_docx):
    p = find(sample_docx, "(1) Purpose")
    runs = {r.text.strip(): r for r in p.runs if r.text.strip()}
    assert runs["(1) Purpose"].bold and runs["(1) Purpose"].highlight == "#FFFF00"
    assert runs["working capital"].color == "#C00000"
    assert runs["rd"].superscript
    assert runs["Tung Hui"].italic


def test_bullets_numbers_and_blank_lines(sample_docx):
    kinds = [(b.kind, b.level, b.text) for b in sample_docx.blocks if b.page == 1]
    i = kinds.index(("bullet", 0, "Top bullet"))
    assert kinds[i - 1][0] == "paragraph" and sample_docx.blocks[i - 1].spacer  # one blank line kept
    assert not sample_docx.blocks[i - 2].spacer
    assert kinds[i + 1] == ("bullet", 1, "Sub bullet")
    # Word's automatic numbering becomes a real Google Docs numbered list
    first, second = sample_docx.blocks[i + 2], sample_docx.blocks[i + 3]
    assert (first.text, second.text) == ("First step", "Second step")
    assert first.kind == "bullet" and first.numbered == "NUMBERED_DECIMAL_ALPHA_ROMAN" and first.level == 0


def test_line_break_starts_new_line(sample_docx):
    texts = [b.text for b in sample_docx.blocks]
    assert "Line one" in texts and "line two" in texts


def test_table_merge_and_shading(sample_docx):
    t = next(b for b in sample_docx.blocks if b.kind == "table").table
    assert (t.n_rows, t.n_cols) == (2, 2)
    head = t.cell(0, 0)
    assert head.colspan == 2 and head.text == "Merged header"
    assert t.cell(1, 0).background == "#9CC2E5"


def test_picture(sample_docx):
    img = next(b for b in sample_docx.blocks if b.kind == "image").image
    assert img.png.startswith(b"\x89PNG") and round(img.width_pt) == 200


def test_old_doc_format_message(tmp_path):
    f = tmp_path / "old.doc"
    f.write_bytes(b"not really")
    with pytest.raises(ParseError, match="Save As"):
        parse_file(f)


# --------------------------------------------------------------------------- #
# Course samples (local only)
# --------------------------------------------------------------------------- #

def test_banking_sample():
    doc = parse_sample(BANKING)
    assert find(doc, "1. CORPORATE FINANCING").style_key == "word_h1"      # "Law Heading 1"
    assert find(doc, "A. Distinction Between Equity").style_key == "word_h2"
    h4 = find(doc, "(a) Mechanics", "heading")
    assert h4.style_key == "word_h4" and h4.runs[0].highlight == "#FFFF00"
    sub = find(doc, "Asset (capital", "bullet")
    assert sub.level == 1 and sub.runs[0].color == "#0070C0"
    assert sum(1 for b in doc.blocks if b.kind == "image") == 2


def test_civ_sample():
    doc = parse_sample(CIV)
    assert find(doc, "Introduction: Jurisdiction and Power").style_key == "word_h1"  # "Law Heading 1"
    third = find(doc, "The warrant serves as proof")
    assert third.kind == "bullet" and third.level == 2
    courts = next(b for b in doc.blocks if b.kind == "table" and "Small Claims Tribunal" in b.text).table
    assert courts.cell(0, 0).rowspan == 2 and courts.cell(0, 1).colspan == 2
    assert courts.cell(0, 0).background == "#FFD966"
