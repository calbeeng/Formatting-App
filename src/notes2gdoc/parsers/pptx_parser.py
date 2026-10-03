"""PowerPoint (.pptx) parser.

A .pptx states most of its structure, so this reads it directly:

* Slide title: the title placeholder if the slide has one. Many decks (e.g.
  ones built from blank slides) use a plain text box instead, so otherwise
  the top-most short text box near the top of the slide is the title.
  Consecutive slides with the same title are merged under one heading, slide
  title colour is dropped, and the first slide (title slide) is unticked, as
  for PDF slide decks.
* Section header layouts become "section_title" headings (Heading 1 by default).
* Bullets and their levels come straight from PowerPoint: a paragraph is a
  bullet if it (or its text box's list style) says so, or if it sits in a
  content placeholder, which shows bullets by default. Non-bullet paragraphs
  at a deeper level are indented.
* Reading order: text boxes side by side are read left column first, then
  right column; otherwise top to bottom.
* Tables become tables (merged cells, column widths and cell colours kept).
* Pictures (e.g. exhibits) become pictures. Tiny logos are skipped.
* Diagrams built from boxes joined by lines/arrows are redrawn as a picture,
  a close approximation of the original (PowerPoint doesn't store a picture
  of them). Charts and SmartArt can't be converted and are reported.
* Slide numbers, footers, "©All rights reserved" lines and "Access the text
  alternative…" accessibility notes are left out.

Page numbers in the outline are slide numbers.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from html import escape
from pathlib import Path

import pymupdf as fitz

from ..config import Settings
from ..lists import bulletise_slide
from ..model import Block, Document, Image, Run, Table, TableCell, merge_runs, normalise_whitespace
from ..numbering import int_to_roman
from .base import ParseError
from .images import to_png
from .symbols import fix_symbols
from .pdf.slides import BOILERPLATE_TITLES, _norm
from .pdf.structure import strip_uniform_style

try:
    import pptx
    from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
    from pptx.oxml.ns import qn
except ImportError:  # pragma: no cover
    pptx = None

EMU_PER_PT = 12700
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

# Text that's never wanted in notes
_JUNK = re.compile(r"access the text alternative|^\s*©|all rights reserved\.?\s*$|^\s*\d{1,3}\s*$", re.I)
# Smallest picture worth keeping (points); smaller ones are icons/logos
MIN_PICTURE_PT = 60
# Titles without a title placeholder: start in the top 18% of the slide
TITLE_ZONE = 0.18


@dataclass
class _Item:
    kind: str                      # "text", "box", "table", "picture", "line", "shape", "other"
    x0: float
    y0: float
    x1: float
    y1: float
    shape: object = None
    ph: object = None              # placeholder type, if a placeholder
    order: int = 0                 # position in the slide's drawing order
    line: tuple | None = None      # ((bx, by), (ex, ey), head_arrow, tail_arrow, colour, width)
    text: str = ""

    @property
    def cx(self):
        return (self.x0 + self.x1) / 2

    @property
    def cy(self):
        return (self.y0 + self.y1) / 2

    @property
    def width(self):
        return self.x1 - self.x0


@dataclass
class _Transform:
    """Maps a shape's EMU coordinates to slide points (handles groups)."""

    ox: float = 0.0
    oy: float = 0.0
    sx: float = 1.0 / EMU_PER_PT
    sy: float = 1.0 / EMU_PER_PT
    cx: float = 0.0   # child offset subtracted first
    cy: float = 0.0

    def point(self, x, y):
        return self.ox + (x - self.cx) * self.sx, self.oy + (y - self.cy) * self.sy

    def child(self, group) -> "_Transform":
        xfrm = group._element.grpSpPr.find(A + "xfrm")
        if xfrm is None:
            return self
        off, ext = xfrm.find(A + "off"), xfrm.find(A + "ext")
        choff, chext = xfrm.find(A + "chOff"), xfrm.find(A + "chExt")
        if None in (off, ext, choff, chext):
            return self
        ox, oy = self.point(int(off.get("x")), int(off.get("y")))
        ecx, ecy = int(ext.get("cx")), int(ext.get("cy"))
        ccx, ccy = int(chext.get("cx")) or 1, int(chext.get("cy")) or 1
        return _Transform(ox, oy, self.sx * ecx / ccx, self.sy * ecy / ccy,
                          int(choff.get("x")), int(choff.get("y")))


# --------------------------------------------------------------------------- #
# Theme colours (for redrawing diagrams)
# --------------------------------------------------------------------------- #

_SCHEME_ALIAS = {"bg1": "lt1", "tx1": "dk1", "bg2": "lt2", "tx2": "dk2"}


def _theme_colours(prs) -> dict[str, str]:
    colours = {"dk1": "#000000", "lt1": "#FFFFFF", "dk2": "#44546A", "lt2": "#E7E6E6",
               "accent1": "#4472C4", "accent2": "#ED7D31", "accent3": "#A5A5A5",
               "accent4": "#FFC000", "accent5": "#5B9BD5", "accent6": "#70AD47"}
    try:
        from pptx.opc.constants import RELATIONSHIP_TYPE as RT
        from lxml import etree

        theme = prs.slide_master.part.part_related_by(RT.THEME)
        root = etree.fromstring(theme.blob)
        scheme = root.find(".//" + A + "clrScheme")
        for el in scheme:
            name = el.tag.replace(A, "")
            srgb = el.find(A + "srgbClr")
            sysc = el.find(A + "sysClr")
            if srgb is not None:
                colours[name] = "#" + srgb.get("val").upper()
            elif sysc is not None and sysc.get("lastClr"):
                colours[name] = "#" + sysc.get("lastClr").upper()
    except Exception:
        pass
    return colours


def _colour_of(el, theme) -> str | None:
    """Colour from an element containing a:srgbClr or a:schemeClr, including
    PowerPoint's lighter/darker adjustments ("Blue, Accent 1, Lighter 80%"
    is stored as accent1 with lumMod/lumOff)."""
    if el is None:
        return None
    clr = el.find(A + "srgbClr")
    base = "#" + clr.get("val").upper() if clr is not None else None
    if clr is None:
        clr = el.find(A + "schemeClr")
        if clr is None:
            return None
        base = theme.get(_SCHEME_ALIAS.get(clr.get("val"), clr.get("val")))
    if base is None:
        return None
    return _apply_modifiers(base, clr)


def _apply_modifiers(hex_: str, clr) -> str:
    import colorsys

    r, g, b = (int(hex_[i:i + 2], 16) / 255 for i in (1, 3, 5))
    for mod in clr:
        tag = mod.tag.replace(A, "")
        val = int(mod.get("val", "100000")) / 100000
        if tag in ("lumMod", "lumOff"):
            h, l, s = colorsys.rgb_to_hls(r, g, b)
            l = l * val if tag == "lumMod" else l + val
            r, g, b = colorsys.hls_to_rgb(h, min(1.0, max(0.0, l)), s)
        elif tag == "tint":    # towards white
            r, g, b = (c + (1 - c) * (1 - val) for c in (r, g, b))
        elif tag == "shade":   # towards black
            r, g, b = (c * val for c in (r, g, b))
    return "#" + "".join(f"{round(c * 255):02X}" for c in (r, g, b))


# --------------------------------------------------------------------------- #
# The parser
# --------------------------------------------------------------------------- #

class PptxParser:
    extensions = (".pptx",)

    def parse(self, path: Path, settings: Settings) -> Document:
        if pptx is None:
            raise ParseError("PowerPoint support isn't installed (python-pptx).")
        try:
            prs = pptx.Presentation(str(path))
        except Exception as exc:
            raise ParseError(f"Couldn't open this PowerPoint file: {exc}") from exc
        self.prs = prs
        self.W = prs.slide_width / EMU_PER_PT
        self.H = prs.slide_height / EMU_PER_PT
        self.theme = _theme_colours(prs)
        self.settings = settings
        self.warnings: list[str] = []

        blocks: list[Block] = []
        prev_title: str | None = None
        prev_title_slide = 0
        n_slides = len(prs.slides)
        for idx, slide in enumerate(prs.slides):
            slide_no = idx + 1
            items = self._collect(slide)
            layout = (slide.slide_layout.name or "").lower()
            title_item = self._find_title(items)
            title_text = _clean_text(title_item.shape.text_frame.text) if title_item else ""
            body_items = [it for it in items if it is not title_item]
            texts = [it for it in body_items if it.kind in ("text", "box")]

            # --- Title slide (first slide): kept, unticked -------------------
            if idx == 0 and ("title slide" in layout or not self._has_bullets(texts)):
                for it in ([title_item] if title_item else []) + sorted(texts, key=lambda i: (i.y0, i.x0)):
                    for b in self._text_blocks(it, slide_no):
                        b.selected = False
                        b.note = "title slide (unticked by default)"
                        blocks.append(b)
                continue

            # --- Section divider slide ----------------------------------------
            if "section" in layout and title_text:
                blocks.append(self._heading(title_item, "section_title", slide_no, "section header slide"))
                prev_title = None
                for it in self._ordered(body_items):
                    blocks.extend(self._item_blocks(it, slide_no))
                continue

            slide_blocks: list[Block] = []
            norm = _norm(title_text)
            boilerplate = norm in BOILERPLATE_TITLES or (
                idx == n_slides - 1 and "all rights reserved" in " ".join(i.text.lower() for i in texts))
            if title_text:
                if norm == prev_title and not boilerplate:
                    merged = f"slide {slide_no}: same title as slide {prev_title_slide}, merged"
                else:
                    slide_blocks.append(self._heading(title_item, "slide_title", slide_no, f"slide {slide_no} title"))
                    prev_title, prev_title_slide, merged = norm, slide_no, ""
            else:
                merged = ""

            # --- Diagrams drawn with boxes and connector lines ----------------
            diagram, body_items = self._diagram(body_items, title_item, slide_no)
            body: list[Block] = []
            for it in self._ordered(body_items + ([diagram] if diagram else [])):
                body.extend(self._item_blocks(it, slide_no))
            if not boilerplate:
                # Slides without bullets: each paragraph becomes a point
                body = bulletise_slide(body, lambda b: b.level * 36.0)
            slide_blocks.extend(body)

            if merged and slide_blocks:
                slide_blocks[0].note = merged + (f"; {slide_blocks[0].note}" if slide_blocks[0].note else "")
            if boilerplate:
                for b in slide_blocks:
                    b.selected = not settings.skip_boilerplate_slides
                    b.note = f"boilerplate slide {slide_no}" + (f"; {b.note}" if b.note else "")
                prev_title = None
            blocks.extend(slide_blocks)

        if not blocks:
            raise ParseError("This PowerPoint file doesn't seem to contain any text.")
        return Document(blocks, str(path), layout="slides", page_count=n_slides, warnings=self.warnings)

    # ------------------------------------------------------------------ #
    # Collecting shapes
    # ------------------------------------------------------------------ #
    def _collect(self, slide) -> list[_Item]:
        items: list[_Item] = []
        counter = [0]

        def visit(shapes, tf: _Transform):
            for sh in shapes:
                counter[0] += 1
                st = sh.shape_type
                if st == MSO_SHAPE_TYPE.GROUP:
                    visit(sh.shapes, tf.child(sh))
                    continue
                ph = None
                if sh.is_placeholder:
                    ph = sh.placeholder_format.type
                    if ph in (PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.FOOTER,
                              PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.HEADER):
                        continue
                # (placeholders without their own position inherit the layout's)
                if sh.left is None or sh.width is None:
                    continue
                x0, y0 = tf.point(sh.left or 0, sh.top or 0)
                x1, y1 = tf.point((sh.left or 0) + (sh.width or 0), (sh.top or 0) + (sh.height or 0))
                item = _Item("other", min(x0, x1), min(y0, y1), max(x0, x1), max(y0, y1), sh, ph, counter[0])

                if getattr(sh, "has_table", False) and sh.has_table:
                    item.kind = "table"
                elif getattr(sh, "has_chart", False) and sh.has_chart:
                    self.warnings.append(f"Slide {slide_no_of(slide, self.prs)}: a chart was left out "
                                         "(charts can't be converted).")
                    continue
                elif st == MSO_SHAPE_TYPE.PICTURE or (ph is not None and hasattr(sh, "image")):
                    item.kind = "picture"
                elif type(sh).__name__ == "Connector" or st == MSO_SHAPE_TYPE.LINE:
                    item.kind = "line"
                    item.line = self._line_geometry(sh, tf)
                elif sh.has_text_frame and sh.text_frame.text.strip():
                    item.text = sh.text_frame.text
                    is_box = st == MSO_SHAPE_TYPE.AUTO_SHAPE and not sh.is_placeholder
                    item.kind = "box" if is_box else "text"
                elif st == MSO_SHAPE_TYPE.AUTO_SHAPE:
                    item.kind = "shape"  # e.g. an arrow or decoration without text
                elif "graphicFrame" in sh._element.tag:
                    # Equations and other embedded objects carry a preview
                    # picture (often WMF, which we can draw); SmartArt doesn't.
                    blip = sh._element.find(".//" + A + "blip")
                    rid = blip.get(qn("r:embed")) if blip is not None else None
                    if rid is None:
                        self.warnings.append(f"Slide {slide_no_of(slide, self.prs)}: a SmartArt graphic was "
                                             "left out (it can't be converted).")
                        continue
                    item.kind = "equation"
                    item.text = slide.part.related_part(rid).blob  # the preview picture
                    item.line = None
                else:
                    continue
                items.append(item)

        visit(slide.shapes, _Transform())
        return items

    def _line_geometry(self, sh, tf):
        try:
            b = tf.point(sh.begin_x, sh.begin_y)
            e = tf.point(sh.end_x, sh.end_y)
        except Exception:
            x0, y0 = tf.point(sh.left, sh.top)
            x1, y1 = tf.point(sh.left + sh.width, sh.top + sh.height)
            b, e = (x0, y0), (x1, y1)
        ln = sh._element.spPr.find(A + "ln")
        head = tail = False
        colour, width = "#404040", 1.0
        if ln is not None:
            he, te = ln.find(A + "headEnd"), ln.find(A + "tailEnd")
            head = he is not None and he.get("type", "none") != "none"
            tail = te is not None and te.get("type", "none") != "none"
            colour = _colour_of(ln.find(A + "solidFill"), self.theme) or colour
            if ln.get("w"):
                width = max(0.75, int(ln.get("w")) / EMU_PER_PT)
        return (b, e, head, tail, colour, width)

    # ------------------------------------------------------------------ #
    # Title, ordering
    # ------------------------------------------------------------------ #
    def _find_title(self, items: list[_Item]) -> _Item | None:
        for it in items:
            if it.ph in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE) \
                    and it.shape.has_text_frame and it.shape.text_frame.text.strip():
                return it
        # No title placeholder: the top-most short, bullet-less text box near the top
        candidates = []
        for it in items:
            if it.kind not in ("text", "box") or it.y0 > self.H * TITLE_ZONE:
                continue
            paras = [p for p in it.shape.text_frame.paragraphs if p.text.strip()]
            text = " ".join(p.text.strip() for p in paras)
            if 0 < len(paras) <= 2 and len(text) <= 150 and not any(self._bullet_kind(p, it) for p in paras):
                candidates.append(it)
        return min(candidates, key=lambda i: (i.y0, i.x0)) if candidates else None

    def _ordered(self, items: list[_Item]) -> list[_Item]:
        """Reading order: left column then right column when text sits side by
        side; otherwise top to bottom, left to right."""
        content = [it for it in items if it.kind in ("text", "box", "table", "picture", "diagram", "equation")]
        if not content:
            return []
        # Only items clearly on one side of the slide define the columns; a
        # centred label spanning both (e.g. "Degree of risk aversion") would
        # otherwise hide the gap between them.
        sided = [it for it in content if it.width < 0.55 * self.W
                 and (it.x1 <= 0.65 * self.W or it.x0 >= 0.35 * self.W)]
        gutter = _find_gutter(sided, self.W)
        if gutter is None:
            return sorted(content, key=lambda i: (round(i.y0 / 12), i.x0))
        left = [i for i in content if i.x1 <= gutter]
        right = [i for i in content if i.x0 >= gutter]
        full = [i for i in content if i not in left and i not in right]
        col_top = min(i.y0 for i in left + right)
        key = lambda i: (i.y0, i.x0)  # noqa: E731
        return (sorted([i for i in full if i.y0 < col_top], key=key) + sorted(left, key=key)
                + sorted(right, key=key) + sorted([i for i in full if i.y0 >= col_top], key=key))

    def _has_bullets(self, items: list[_Item]) -> bool:
        return any(self._bullet_kind(p, it) for it in items for p in it.shape.text_frame.paragraphs if p.text.strip())

    # ------------------------------------------------------------------ #
    # Turning items into blocks
    # ------------------------------------------------------------------ #
    def _item_blocks(self, it: _Item, slide_no: int) -> list[Block]:
        if it.kind in ("text", "box"):
            return self._text_blocks(it, slide_no)
        if it.kind == "table":
            return [self._table(it, slide_no)]
        if it.kind == "picture":
            return self._picture(it, slide_no)
        if it.kind == "equation":
            png = to_png(it.text)
            if png is None:
                self.warnings.append(f"Slide {slide_no}: an embedded object's picture couldn't be converted.")
                return []
            w, h = it.x1 - it.x0, it.y1 - it.y0
            return [Block("image", image=Image(png, w, h, f"Equation from slide {slide_no}"), page=slide_no,
                          note="equation / embedded object (drawn from its preview)")]
        if it.kind == "diagram":
            return [it.shape]  # already a Block
        return []

    def _heading(self, item: _Item, style_key: str, slide_no: int, note: str) -> Block:
        runs: list[Run] = []
        for i, p in enumerate(q for q in item.shape.text_frame.paragraphs if q.text.strip()):
            if i:
                runs.append(Run(" "))
            runs.extend(self._runs(p))
        runs = [Run(r.text, r.bold, r.italic, r.underline, r.superscript, r.subscript) for r in runs]  # no colour
        runs = strip_uniform_style(merge_runs(runs), bold=True, italic=True)
        b = Block("heading", normalise_whitespace(runs), style_key=style_key, page=slide_no, note=note)
        return b

    def _text_blocks(self, it: _Item, slide_no: int) -> list[Block]:
        out: list[Block] = []
        counters: dict[int, int] = {}
        anchor: Block | None = None  # last line not indented with leading spaces
        for p in it.shape.text_frame.paragraphs:
            text = p.text.strip()
            if not text or _JUNK.search(text):
                continue
            runs = normalise_whitespace(self._runs(p))
            if not runs:
                continue
            kind = self._bullet_kind(p, it)
            level = p.level
            # Plain lines indented by typing spaces at the start (or by a left
            # margin) sit under the line above: "   = E(Ri) + …" under its
            # bullet, "      βi GDP = …" under "where …".
            # Each such line is indented one step past the last line that
            # WASN'T indented this way (its "anchor"), so a run of them lines up.
            lead = len(p.text) - len(p.text.lstrip(" 	"))
            if not kind:
                marl = int(p._p.pPr.get("marL", "0")) if p._p.pPr is not None else 0
                if anchor is not None and (lead >= 2 and anchor.kind == "bullet" or lead >= 4):
                    level = max(level, anchor.level + 1)
                    spaced = True
                else:
                    spaced = False
                    if marl >= 228600:  # a left margin of at least a quarter inch
                        level = max(level, round(marl / 457200))
            else:
                spaced = False
            if isinstance(kind, tuple):  # automatic numbering ("1.", "a)")
                scheme, start = kind
                counters[level] = counters.get(level, start - 1) + 1
                for deeper in [k for k in counters if k > level]:
                    del counters[deeper]
                label = _auto_number(scheme, counters[level])
                out.append(Block("paragraph", [Run(label + " ")] + runs, level=level + 1, page=slide_no,
                                 note=f"numbered item {label!r}"))
            elif kind:
                out.append(Block("bullet", runs, level=level, page=slide_no, note=f"bullet, level {level + 1}"))
            else:
                note = "text in a shape" if it.kind == "box" else ""
                if spaced:
                    note = "line indented with spaces"
                out.append(Block("paragraph", runs, level=level, page=slide_no, note=note))
            if not spaced:
                anchor = out[-1]
        return out

    def _bullet_kind(self, p, it: _Item):
        """True for a bullet, ("scheme", start) for automatic numbering, or
        False for plain text."""
        lvl = f"{A}lvl{p.level + 1}pPr"
        sources = [p._p.pPr]
        lst = it.shape.text_frame._txBody.find(A + "lstStyle")
        if lst is not None:
            sources.append(lst.find(lvl))
        sources += self._inherited_styles(it, lvl)
        for src in sources:
            if src is None:
                continue
            if src.find(A + "buNone") is not None:
                return False
            auto = src.find(A + "buAutoNum")
            if auto is not None:
                return (auto.get("type", "arabicPeriod"), int(auto.get("startAt", "1")))
            if src.find(A + "buChar") is not None or src.find(A + "buBlip") is not None:
                return True
        # Content placeholders show bullets unless told otherwise
        return it.ph in (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT) if it.ph is not None else False

    def _inherited_styles(self, it: _Item, lvl: str) -> list:
        """A placeholder's paragraph style for this level from the slide
        layout, then from the slide master's body text style. (Some templates
        switch bullets off for the top level there: "Exhibit 4.2" is plain
        text, only the level below has bullets.)"""
        sh = it.shape
        if it.ph is None or not getattr(sh, "is_placeholder", False):
            return []
        out = []
        try:
            layout = sh.part.slide.slide_layout
            idx = sh.placeholder_format.idx
            lp = next((x for x in layout.placeholders if x.placeholder_format.idx == idx), None)
            if lp is not None:
                lst = lp._element.find(".//" + A + "lstStyle")
                if lst is not None:
                    out.append(lst.find(lvl))
            if it.ph in (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT):
                body = layout.slide_master._element.find(".//" + qn("p:bodyStyle"))
                if body is not None:
                    out.append(body.find(lvl))
        except Exception:
            return out
        return out

    def _runs(self, p) -> list[Run]:
        runs = []
        for r in p.runs:
            if not r.text:
                continue
            f = r.font
            rPr = r._r.find(A + "rPr")
            baseline = int(rPr.get("baseline", "0")) if rPr is not None else 0
            colour = None
            try:
                if f.color is not None and f.color.type is not None and str(f.color.type).startswith("RGB"):
                    colour = "#" + str(f.color.rgb).upper()
            except Exception:
                colour = None
            if colour and max(int(colour[1:3], 16), int(colour[3:5], 16), int(colour[5:7], 16)) < 0x30:
                colour = None
            highlight = _colour_of(rPr.find(A + "highlight"), self.theme) if rPr is not None else None
            # The run's font, and its separate font for symbol characters
            font = sym = None
            if rPr is not None:
                latin, sym_el = rPr.find(A + "latin"), rPr.find(A + "sym")
                font = latin.get("typeface") if latin is not None else None
                sym = sym_el.get("typeface") if sym_el is not None else None
            runs.append(Run(
                fix_symbols(r.text.replace("\x0b", " "), font, sym),
                bold=bool(f.bold), italic=bool(f.italic),
                underline=bool(f.underline) and f.underline is not False,
                superscript=baseline > 0, subscript=baseline < 0,
                color=colour, highlight=highlight,
            ))
        return runs

    def _table(self, it: _Item, slide_no: int) -> Block:
        tbl = it.shape.table
        n_rows, n_cols = len(tbl.rows), len(tbl.columns)
        cells = []
        for r in range(n_rows):
            for c in range(n_cols):
                cell = tbl.cell(r, c)
                if cell.is_spanned:
                    continue
                fake = _Item("text", 0, 0, 0, 0, cell, None)
                blocks = []
                for p in cell.text_frame.paragraphs:
                    text = p.text.strip()
                    if not text:
                        continue
                    kind = self._bullet_kind(p, fake)
                    runs = normalise_whitespace(self._runs(p))
                    if runs:
                        blocks.append(Block("bullet" if kind is True else "paragraph", runs, level=p.level,
                                            page=slide_no))
                bg = None
                try:
                    if cell.fill.type is not None and str(cell.fill.fore_color.type).startswith("RGB"):
                        bg = "#" + str(cell.fill.fore_color.rgb).upper()
                except Exception:
                    bg = None
                if bg is None:
                    tcPr = cell._tc.tcPr
                    bg = _colour_of(tcPr.find(A + "solidFill"), self.theme) if tcPr is not None else None
                cells.append(TableCell(r, c, blocks, rowspan=cell.span_height if cell.is_merge_origin else 1,
                                       colspan=cell.span_width if cell.is_merge_origin else 1,
                                       background=bg if bg not in ("#FFFFFF",) else None))
        widths = [col.width / EMU_PER_PT for col in tbl.columns]
        return Block("table", table=Table(n_rows, n_cols, cells, col_widths=widths), page=slide_no,
                     note=f"table {n_rows} rows x {n_cols} columns")

    def _picture(self, it: _Item, slide_no: int) -> list[Block]:
        w, h = it.x1 - it.x0, it.y1 - it.y0
        if w < MIN_PICTURE_PT and h < MIN_PICTURE_PT:
            return []  # icon/logo
        try:
            blob = it.shape.image.blob
        except Exception:
            return []
        png = to_png(blob)
        if png is None:
            self.warnings.append(f"Slide {slide_no}: a picture in a format Google Docs can't use was skipped.")
            return []
        return [Block("image", image=Image(png, w, h, f"Picture from slide {slide_no}"), page=slide_no,
                      note=f"picture {w:.0f}x{h:.0f}pt")]

    # ------------------------------------------------------------------ #
    # Diagrams
    # ------------------------------------------------------------------ #
    def _diagram(self, items: list[_Item], title: _Item | None, slide_no: int):
        """If boxes on the slide are joined by lines/arrows, redraw that area as
        a picture. Returns (diagram item or None, remaining items)."""
        boxes = [i for i in items if i.kind == "box"]
        lines = [i for i in items if i.kind == "line"]
        if title is not None:  # ignore the decorative line under the title
            lines = [ln for ln in lines if not (abs(ln.cy - title.y1) < 20 and ln.y1 - ln.y0 < 4)]
        if len(boxes) < 2 or not lines:
            return None, items
        rx0, ry0 = min(b.x0 for b in boxes), min(b.y0 for b in boxes)
        rx1, ry1 = max(b.x1 for b in boxes), max(b.y1 for b in boxes)
        # A connector must join two different boxes (each end within 50pt of
        # a box). A line under a boxed title only touches one box.
        def near(pt, box):
            return box.x0 - 50 <= pt[0] <= box.x1 + 50 and box.y0 - 50 <= pt[1] <= box.y1 + 50

        def joins(ln):
            (b, e) = ln.line[0], ln.line[1]
            ends_b = [x for x in boxes if near(b, x)]
            ends_e = [x for x in boxes if near(e, x)]
            return any(x is not y for x in ends_b for y in ends_e)

        connecting = [ln for ln in lines if joins(ln)]
        if not connecting:
            return None, items
        region = [rx0, ry0, rx1, ry1]
        for ln in connecting:
            region = [min(region[0], ln.x0), min(region[1], ln.y0), max(region[2], ln.x1), max(region[3], ln.y1)]
        members = [i for i in items if i.kind in ("box", "shape", "line", "picture", "text")
                   and region[0] - 2 <= i.cx <= region[2] + 2 and region[1] - 2 <= i.cy <= region[3] + 2
                   and (i.kind != "text" or len(i.text) < 120)]
        # Short labels just beside the diagram (e.g. "Size Factor" in a dashed
        # box to the left of a row of boxes) belong to it too
        for i in items:
            if i in members or i.kind not in ("text", "box") or len(i.text) >= 60:
                continue
            beside = (region[1] <= i.cy <= region[3]) and (region[0] - 150 <= i.x1 <= region[0] + 5
                                                           or region[2] - 5 <= i.x0 <= region[2] + 150)
            if beside:
                members.append(i)
        for m in members:
            region = [min(region[0], m.x0), min(region[1], m.y0), max(region[2], m.x1), max(region[3], m.y1)]
        png = self._render(members, region)
        alt = "Diagram: " + "; ".join(_clean_text(m.text) for m in members if m.text.strip())[:1000]
        w, h = region[2] - region[0], region[3] - region[1]
        block = Block("image", image=Image(png, w + 12, h + 12, alt), page=slide_no,
                      note="diagram redrawn from the slide's shapes")
        item = _Item("diagram", *region, shape=block)
        return item, [i for i in items if i not in members]

    def _render(self, members: list[_Item], region) -> bytes:
        pad = 6
        x0, y0 = region[0] - pad, region[1] - pad
        doc = fitz.open()
        page = doc.new_page(width=region[2] - region[0] + 2 * pad, height=region[3] - region[1] + 2 * pad)
        for m in sorted(members, key=lambda i: i.order):
            r = fitz.Rect(m.x0 - x0, m.y0 - y0, m.x1 - x0, m.y1 - y0)
            if m.kind == "line":
                (bx, by), (ex, ey), head, tail, colour, width = m.line
                p1, p2 = fitz.Point(bx - x0, by - y0), fitz.Point(ex - x0, ey - y0)
                col = _rgb(colour)
                page.draw_line(p1, p2, color=col, width=width)
                if tail:
                    _arrowhead(page, p1, p2, col, width)
                if head:
                    _arrowhead(page, p2, p1, col, width)
            elif m.kind == "picture":
                try:
                    page.insert_image(r, stream=m.shape.image.blob)
                except Exception:
                    pass
            elif m.kind in ("box", "shape"):
                fill, stroke, text_colour = self._shape_colours(m.shape)
                geom = m.shape._element.spPr.find(A + "prstGeom")
                prst = geom.get("prst") if geom is not None else "rect"
                if prst in ("ellipse", "flowChartConnector"):
                    page.draw_oval(r, color=stroke, fill=fill, width=1)
                elif prst in ("roundRect", "flowChartAlternateProcess"):
                    page.draw_rect(r, color=stroke, fill=fill, width=1, radius=0.15)
                elif "Arrow" in prst and not m.text.strip():
                    _block_arrow(page, r, prst, fill or stroke)
                else:
                    ln = m.shape._element.spPr.find(A + "ln")
                    dash = ln.find(A + "prstDash") if ln is not None else None
                    dashes = "[3 2] 0" if dash is not None and dash.get("val", "solid") != "solid" else None
                    page.draw_rect(r, color=stroke, fill=fill, width=1, dashes=dashes)
                if m.text.strip():
                    self._draw_text(page, r, m.shape, text_colour, centred=True)
            elif m.kind == "text":
                ln = m.shape._element.spPr.find(A + "ln")
                outline = _colour_of(ln.find(A + "solidFill"), self.theme) if ln is not None else None
                if outline:
                    dash = ln.find(A + "prstDash")
                    dashes = "[3 2] 0" if dash is not None and dash.get("val", "solid") != "solid" else None
                    page.draw_rect(r, color=_rgb(outline), width=0.75, dashes=dashes)
                self._draw_text(page, r, m.shape, (0, 0, 0), centred=bool(outline))
        png = page.get_pixmap(dpi=144).tobytes("png")
        doc.close()
        return png

    def _shape_colours(self, sh):
        """(fill, outline, text colour) as RGB tuples, resolving theme colours."""
        spPr = sh._element.spPr
        style = sh._element.find(qn("p:style"))
        fill = stroke = None
        if spPr.find(A + "noFill") is not None:
            fill = None
        else:
            fill = _colour_of(spPr.find(A + "solidFill"), self.theme)
            if fill is None and style is not None:
                ref = style.find(A + "fillRef")
                if ref is not None and ref.get("idx", "0") != "0":
                    fill = _colour_of(ref, self.theme)
        ln = spPr.find(A + "ln")
        if ln is not None and ln.find(A + "noFill") is not None:
            stroke = None
        else:
            stroke = _colour_of(ln.find(A + "solidFill"), self.theme) if ln is not None else None
            if stroke is None and style is not None:
                ref = style.find(A + "lnRef")
                if ref is not None and ref.get("idx", "0") != "0":
                    stroke = _colour_of(ref, self.theme)
        text = None
        if style is not None:
            text = _colour_of(style.find(A + "fontRef"), self.theme)
        # Make sure the text is readable on what's behind it (the shape's fill,
        # or the white page when the shape has none)
        behind = fill or "#FFFFFF"
        if text is None or _dark(text) == _dark(behind):
            text = "#FFFFFF" if _dark(behind) else "#000000"
        return _rgb(fill), _rgb(stroke), _rgb(text)

    def _draw_text(self, page, rect, sh, colour, centred: bool):
        paras = []
        size = 12.0
        for p in sh.text_frame.paragraphs:
            if not p.text.strip():
                continue
            parts = []
            for r in p.runs:
                t = escape(r.text)
                if r.font.size:
                    size = r.font.size.pt
                if r.font.bold:
                    t = f"<b>{t}</b>"
                if r.font.italic:
                    t = f"<i>{t}</i>"
                parts.append(t)
            paras.append("".join(parts))
        if not paras:
            return
        r, g, b = (int(c * 255) for c in colour)
        family = "serif" if _serif_font(sh) else "sans-serif"
        css = (f"* {{font-family: {family}; font-size: {size:.0f}px; color: rgb({r},{g},{b});"
               f" text-align: {'center' if centred else 'left'}; margin: 0;}}")
        html = "".join(f"<p>{p}</p>" for p in paras)
        inner = fitz.Rect(rect.x0 + 3, rect.y0 + 2, rect.x1 - 3, rect.y1 - 2)
        if centred:
            # Measure on a scratch page, then centre vertically
            scratch = fitz.open()
            sp = scratch.new_page(width=page.rect.width, height=page.rect.height + 500)
            spare, _ = sp.insert_htmlbox(fitz.Rect(inner.x0, 0, inner.x1, inner.height), html, css=css)
            scratch.close()
            if spare > 0:
                inner = fitz.Rect(inner.x0, inner.y0 + spare / 2, inner.x1, inner.y1)
        page.insert_htmlbox(inner, html, css=css, scale_low=0.5)


# --------------------------------------------------------------------------- #
# Small helpers
# --------------------------------------------------------------------------- #

def slide_no_of(slide, prs) -> int:
    for i, s in enumerate(prs.slides, 1):
        if s.slide_id == slide.slide_id:
            return i
    return 0


def _serif_font(sh) -> bool:
    """Does the shape's text use a serif font like Times New Roman?"""
    for latin in sh._element.iter(A + "latin"):
        face = (latin.get("typeface") or "").lower()
        if face and not face.startswith("+"):
            return any(k in face for k in ("times", "georgia", "garamond", "cambria", "serif", "book"))
    return False


def _clean_text(text: str) -> str:
    return re.sub(r"\s+", " ", text.replace("\x0b", " ")).strip()


def _find_gutter(items: list[_Item], width: float) -> float | None:
    if len(items) < 2:
        return None
    spans = sorted((i.x0, i.x1) for i in items)
    merged: list[list[float]] = []
    for a, b in spans:
        if merged and a <= merged[-1][1]:
            merged[-1][1] = max(merged[-1][1], b)
        else:
            merged.append([a, b])
    for (_, a1), (b0, _) in zip(merged, merged[1:]):
        mid = (a1 + b0) / 2
        if b0 - a1 >= 6 and 0.3 * width <= mid <= 0.7 * width:
            left = [i for i in items if i.x1 <= mid]
            right = [i for i in items if i.x0 >= mid]
            overlap = min(max(i.y1 for i in left), max(i.y1 for i in right)) - max(
                min(i.y0 for i in left), min(i.y0 for i in right))
            if left and right and overlap > 0:
                return mid
    return None


def _auto_number(scheme: str, n: int) -> str:
    if scheme.startswith("alphaLc"):
        core = chr(96 + (n - 1) % 26 + 1)
    elif scheme.startswith("alphaUc"):
        core = chr(64 + (n - 1) % 26 + 1)
    elif scheme.startswith("romanLc"):
        core = int_to_roman(n)
    elif scheme.startswith("romanUc"):
        core = int_to_roman(n).upper()
    else:
        core = str(n)
    if scheme.endswith("ParenBoth"):
        return f"({core})"
    if scheme.endswith("ParenR"):
        return f"{core})"
    if scheme.endswith("Plain"):
        return core
    return f"{core}."


def _rgb(hex_: str | None):
    if not hex_:
        return None
    return tuple(int(hex_[i:i + 2], 16) / 255 for i in (1, 3, 5))


def _dark(hex_: str) -> bool:
    r, g, b = (int(hex_[i:i + 2], 16) for i in (1, 3, 5))
    return 0.299 * r + 0.587 * g + 0.114 * b < 140


def _arrowhead(page, start, tip, colour, width):
    """A filled triangle at `tip`, pointing away from `start`."""
    import math

    dx, dy = tip.x - start.x, tip.y - start.y
    length = math.hypot(dx, dy) or 1
    ux, uy = dx / length, dy / length
    size = 6 + width * 1.5
    base = fitz.Point(tip.x - ux * size, tip.y - uy * size)
    left = fitz.Point(base.x - uy * size / 2, base.y + ux * size / 2)
    right = fitz.Point(base.x + uy * size / 2, base.y - ux * size / 2)
    page.draw_polyline([tip, left, right, tip], color=colour, fill=colour, width=0.5, closePath=True)


def _block_arrow(page, r, prst, colour):
    """Simple block arrows (rightArrow, downArrow, …) used as connectors."""
    c = colour or (0.27, 0.45, 0.77)
    if prst.startswith(("left", "right")):
        mid = r.y0 + r.height / 2
        body = fitz.Rect(r.x0, mid - r.height / 4, r.x1 - r.height / 2, mid + r.height / 4)
        if prst.startswith("left"):
            body = fitz.Rect(r.x0 + r.height / 2, mid - r.height / 4, r.x1, mid + r.height / 4)
            tip = [fitz.Point(r.x0, mid), fitz.Point(body.x0, r.y0), fitz.Point(body.x0, r.y1)]
        else:
            tip = [fitz.Point(r.x1, mid), fitz.Point(body.x1, r.y0), fitz.Point(body.x1, r.y1)]
    else:
        mid = r.x0 + r.width / 2
        body = fitz.Rect(mid - r.width / 4, r.y0, mid + r.width / 4, r.y1 - r.width / 2)
        if prst.startswith("up"):
            body = fitz.Rect(mid - r.width / 4, r.y0 + r.width / 2, mid + r.width / 4, r.y1)
            tip = [fitz.Point(mid, r.y0), fitz.Point(r.x0, body.y0), fitz.Point(r.x1, body.y0)]
        else:
            tip = [fitz.Point(mid, r.y1), fitz.Point(r.x0, body.y1), fitz.Point(r.x1, body.y1)]
    page.draw_rect(body, color=None, fill=c)
    page.draw_polyline(tip + [tip[0]], color=c, fill=c, closePath=True)
